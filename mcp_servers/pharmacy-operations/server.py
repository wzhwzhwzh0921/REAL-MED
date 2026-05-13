#!/usr/bin/env python3
"""Pharmacy-operations MCP server.

Data model: each eval question maps to one small Excel workbook under
`data/questions/<excel_file>` (single sheet `库存清单`, one drug per row).
A session loads exactly that question's workbook on `initialize_session`,
so the drug universe is naturally scoped per question — no global merge.

`data/question_profiles.jsonl` maps question_id -> {path, excel_file,
query_budget, recommended_tools}.
"""

import json
import os
import sys
import uuid
from pathlib import Path

import openpyxl


MCP_DIR = Path(__file__).parent
# Per-question pharmacy data lives under the eval registry so all bulky
# attachments stay in one place. Override with PHARMACY_DATA_DIR if needed.
_DEFAULT_DATA_DIR = MCP_DIR.parent.parent / "evals" / "registry" / "mcp_data" / "pharmacy-operations"
DATA_DIR = Path(os.environ.get("PHARMACY_DATA_DIR", str(_DEFAULT_DATA_DIR)))
QUESTIONS_DIR = DATA_DIR / "questions"
OUTPUT_DIR = Path(os.environ.get("OUTPUTS_DIR", str(MCP_DIR / "output")))
LOG_DIR = OUTPUT_DIR / "session_logs"
SCHEMA_PATH = MCP_DIR / "tool_schema.json"

SHEET_NAME = "库存清单"

# Column groups (the Excel header uses these Chinese column names verbatim).
IDENTITY_COLS = ["通用名", "品牌名", "剂型", "规格", "包装规格"]
INVENTORY_COLS = ["库存量", "单位", "补货警戒线", "库位", "单位价格(元)"]
BATCH_COLS = ["批号", "效期至"]
CLINICAL_COLS = ["医保类型", "适应症", "用法用量", "禁忌", "儿童用药", "老年用药", "孕妇用药", "不良反应"]
REIMBURSEMENT_COLS = ["医保类型", "适应症", "禁忌", "儿童用药", "老年用药", "孕妇用药"]


def load_jsonl(path: Path, key: str):
    items = {}
    if not path.exists():
        return items
    with path.open("r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            row = json.loads(line)
            items[row[key]] = row
    return items


QUESTION_PROFILES = load_jsonl(DATA_DIR / "question_profiles.jsonl", "question_id")
TOOL_SCHEMA = json.loads(SCHEMA_PATH.read_text(encoding="utf-8"))

LOG_DIR.mkdir(parents=True, exist_ok=True)
SESSIONS_FILE = LOG_DIR / "sessions.json"


# --------------------------------------------------------------------------
# Per-question drug data (loaded from the small Excel workbooks, cached)
# --------------------------------------------------------------------------

_DRUGS_CACHE = {}  # question_id -> list[dict]  (each dict has "drug_id" + Excel columns)


def load_question_drugs(question_id: str):
    if question_id in _DRUGS_CACHE:
        return _DRUGS_CACHE[question_id]
    profile = QUESTION_PROFILES.get(question_id)
    if not profile:
        raise ValueError(f"unknown question_id: {question_id}")
    excel_path = QUESTIONS_DIR / profile["excel_file"]
    if not excel_path.exists():
        raise ValueError(f"missing question workbook: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        _DRUGS_CACHE[question_id] = []
        return []
    header = [str(h) if h is not None else "" for h in rows[0]]
    drugs = []
    for i, raw in enumerate(rows[1:], start=1):
        record = {}
        for col, val in zip(header, raw):
            if col:
                record[col] = "" if val is None else (val if isinstance(val, (int, float)) else str(val))
        record["drug_id"] = f"{question_id}_r{i}"
        record["question_id"] = question_id
        drugs.append(record)
    _DRUGS_CACHE[question_id] = drugs
    return drugs


def drugs_for_session(session):
    return load_question_drugs(session["question_id"])


def drug_index(session):
    return {d["drug_id"]: d for d in drugs_for_session(session)}


# --------------------------------------------------------------------------
# Sessions
# --------------------------------------------------------------------------

def _load_sessions() -> dict:
    if SESSIONS_FILE.exists():
        try:
            return json.loads(SESSIONS_FILE.read_text(encoding="utf-8"))
        except Exception:
            pass
    return {}


def _save_sessions(sessions: dict):
    SESSIONS_FILE.write_text(json.dumps(sessions, ensure_ascii=False, default=str), encoding="utf-8")


def consume_budget(session_id: str, amount: int = 1):
    sessions = _load_sessions()
    session = sessions.get(session_id)
    if not session:
        raise ValueError("invalid session_id")
    session["remaining_budget"] = max(0, session["remaining_budget"] - amount)
    sessions[session_id] = session
    _save_sessions(sessions)
    return session


def get_session(session_id: str):
    sessions = _load_sessions()
    session = sessions.get(session_id)
    if not session:
        raise ValueError("invalid session_id")
    return session


def result_with_session(session, payload: dict):
    return {
        "session_id": session["session_id"],
        "question_id": session["question_id"],
        "remaining_budget": session["remaining_budget"],
        **payload,
    }


def ensure_drug(session, drug_id: str):
    drug = drug_index(session).get(drug_id)
    if not drug:
        raise ValueError("unknown drug_id (not in this question's scope)")
    return drug


# --------------------------------------------------------------------------
# Text search helpers
# --------------------------------------------------------------------------

def text_fields(row):
    return " ".join(str(v) for v in row.values() if v is not None)


def compact_text(value: str) -> str:
    text = (value or "").lower()
    for ch in [" ", "\n", "\t", "：", ":", "（", "）", "(", ")", "，", ",", "。", ".", "/", "-", "_"]:
        text = text.replace(ch, "")
    return text


def query_match(text: str, query: str) -> bool:
    q = (query or "").strip().lower()
    if not q:
        return True
    if q in text:
        return True
    tokens = [tok for tok in q.split() if tok]
    if tokens and all(tok in text for tok in tokens):
        return True
    compact_query = compact_text(q)
    compact_source = compact_text(text)
    if compact_query and compact_query in compact_source:
        return True
    compact_tokens = [compact_text(tok) for tok in tokens if compact_text(tok)]
    if compact_tokens and all(tok in compact_source for tok in compact_tokens):
        return True
    return False


def as_number(value):
    try:
        return float(value)
    except Exception:
        return None


# --------------------------------------------------------------------------
# Records / events
# --------------------------------------------------------------------------

def append_event(session, event_type: str, payload: dict):
    event = {
        "event_id": f"evt_{uuid.uuid4().hex[:12]}",
        "event_type": event_type,
        "session_id": session["session_id"],
        "question_id": session["question_id"],
        "payload": payload,
    }
    session["events"].append(event)
    with Path(session["log_path"]).open("a", encoding="utf-8") as f:
        f.write(json.dumps(event, ensure_ascii=False) + "\n")
    return event


def next_record_id(prefix: str):
    return f"{prefix}_{uuid.uuid4().hex[:10]}"


def _persist_session(session):
    sessions = _load_sessions()
    sessions[session["session_id"]] = session
    _save_sessions(sessions)


def create_record(session, record_type: str, payload: dict):
    record_id = next_record_id(record_type)
    record = {
        "record_id": record_id,
        "record_type": record_type,
        "status": "created",
        **payload,
    }
    session["records"][record_id] = record
    append_event(session, f"create_{record_type}", record)
    _persist_session(session)
    return record


def base_inventory_value(drug):
    return as_number(drug.get("库存量"))


def current_inventory_value(session, drug):
    cur = session["inventory_state"].get(drug["drug_id"])
    if cur is None:
        return base_inventory_value(drug)
    return cur


def snapshot_inventory(session, drug):
    inv = {c: drug.get(c, "") for c in INVENTORY_COLS}
    inv["库存量原始值"] = base_inventory_value(drug)
    inv["会话内库存量"] = current_inventory_value(session, drug)
    return inv


# --------------------------------------------------------------------------
# Session lifecycle
# --------------------------------------------------------------------------

def initialize_question_session(question_id: str, budget_override=None):
    profile = QUESTION_PROFILES.get(question_id)
    if not profile:
        raise ValueError(f"unknown question_id: {question_id}")
    # Eagerly load (and validate) the workbook so failures surface here.
    drugs = load_question_drugs(question_id)
    budget = int(budget_override or profile["query_budget"])
    session_id = f"sess_{uuid.uuid4().hex[:10]}"
    log_path = LOG_DIR / f"{session_id}.jsonl"
    sessions = _load_sessions()
    sessions[session_id] = {
        "session_id": session_id,
        "question_id": question_id,
        "path": profile["path"],
        "remaining_budget": budget,
        "query_budget": budget,
        "events": [],
        "records": {},
        "inventory_state": {},
        "log_path": str(log_path),
    }
    _save_sessions(sessions)
    return {
        "session_id": session_id,
        "question_id": question_id,
        "path": profile["path"],
        "drug_count": len(drugs),
        "query_budget": budget,
        "remaining_budget": budget,
        "log_path": str(log_path),
    }


# --------------------------------------------------------------------------
# Tools
# --------------------------------------------------------------------------

def tool_create_session(args):
    return initialize_question_session(args["question_id"], args.get("budget_override"))


def tool_list_question_profiles(args):
    rows = list(QUESTION_PROFILES.values())
    qid = args.get("question_id")
    path = args.get("path")
    if qid:
        rows = [r for r in rows if r["question_id"] == qid]
    if path:
        rows = [r for r in rows if r["path"] == path]
    return {"items": rows[:100], "count": len(rows)}


def tool_search_drug_catalog(args):
    session = consume_budget(args["session_id"])
    query = (args.get("query") or "").strip().lower()
    filters = args.get("filters") or {}
    limit = int(args.get("limit") or 5)
    hits = []
    for drug in drugs_for_session(session):
        text = text_fields(drug).lower()
        if not query_match(text, query):
            continue
        ok_filters = True
        for k, v in filters.items():
            if str(drug.get(k, "")).lower() != str(v).lower():
                ok_filters = False
                break
        if not ok_filters:
            continue
        hits.append({
            "drug_id": drug["drug_id"],
            "question_id": drug["question_id"],
            "通用名": drug.get("通用名", ""),
            "品牌名": drug.get("品牌名", ""),
            "剂型": drug.get("剂型", ""),
            "规格": drug.get("规格", ""),
            "单位": drug.get("单位", ""),
            "医保类型": drug.get("医保类型", ""),
        })
        if len(hits) >= limit:
            break
    return result_with_session(session, {"items": hits})


def tool_get_drug_base_info(args):
    session = consume_budget(args["session_id"])
    drug = ensure_drug(session, args["drug_id"])
    return result_with_session(session, {
        "drug_id": drug["drug_id"],
        "drug_base_info": {
            "drug_id": drug["drug_id"],
            "question_id": drug["question_id"],
            **{c: drug.get(c, "") for c in IDENTITY_COLS},
        },
    })


def tool_get_inventory_record(args):
    session = consume_budget(args["session_id"])
    drug = ensure_drug(session, args["drug_id"])
    return result_with_session(session, {
        "drug_id": drug["drug_id"],
        "inventory_record": snapshot_inventory(session, drug),
    })


def tool_get_batch_record(args):
    session = consume_budget(args["session_id"])
    drug = ensure_drug(session, args["drug_id"])
    return result_with_session(session, {
        "drug_id": drug["drug_id"],
        "batch_record": {c: drug.get(c, "") for c in BATCH_COLS},
    })


def tool_get_clinical_rule(args):
    session = consume_budget(args["session_id"])
    drug = ensure_drug(session, args["drug_id"])
    return result_with_session(session, {
        "drug_id": drug["drug_id"],
        "clinical_rule": {c: drug.get(c, "") for c in CLINICAL_COLS},
    })


def tool_get_reimbursement_rule(args):
    session = consume_budget(args["session_id"])
    drug = ensure_drug(session, args["drug_id"])
    return result_with_session(session, {
        "drug_id": drug["drug_id"],
        "reimbursement_rule": {
            "drug_id": drug["drug_id"],
            **{c: drug.get(c, "") for c in REIMBURSEMENT_COLS},
        },
    })


def tool_create_dispense_record(args):
    session = get_session(args["session_id"])
    drug = ensure_drug(session, args["drug_id"])
    record = create_record(session, "dispense_record", {
        "drug_id": drug["drug_id"],
        "quantity": args["quantity"],
        "unit": args.get("unit") or drug.get("单位", ""),
        "batch_no": args.get("batch_no", ""),
        "note": args.get("note", ""),
    })
    return result_with_session(session, {"record": record, **record})


def tool_create_inventory_warning(args):
    session = get_session(args["session_id"])
    drug = ensure_drug(session, args["drug_id"])
    record = create_record(session, "inventory_warning", {
        "drug_id": drug["drug_id"],
        "warning_type": args["warning_type"],
        "severity": args.get("severity", "medium"),
        "note": args.get("note", ""),
    })
    return result_with_session(session, {"record": record, **record})


def tool_create_audit_note(args):
    session = get_session(args["session_id"])
    drug_id = args.get("drug_id", "")
    if drug_id:
        ensure_drug(session, drug_id)
    record = create_record(session, "audit_note", {
        "drug_id": drug_id,
        "decision": args["decision"],
        "reason": args.get("reason", ""),
        "note": args.get("note", ""),
    })
    return result_with_session(session, {"record": record, **record})


def tool_create_insurance_intercept(args):
    session = get_session(args["session_id"])
    drug_id = args.get("drug_id", "")
    if drug_id:
        ensure_drug(session, drug_id)
    record = create_record(session, "insurance_intercept", {
        "drug_id": drug_id,
        "decision": args["decision"],
        "reason": args["reason"],
        "note": args.get("note", ""),
    })
    return result_with_session(session, {"record": record, **record})


def tool_create_handoff_note(args):
    session = get_session(args["session_id"])
    drug_id = args.get("drug_id", "")
    if drug_id:
        ensure_drug(session, drug_id)
    record = create_record(session, "handoff_note", {
        "drug_id": drug_id,
        "shift": args.get("shift", ""),
        "note": args["note"],
    })
    return result_with_session(session, {"record": record, **record})


def tool_update_inventory_balance(args):
    session = get_session(args["session_id"])
    drug = ensure_drug(session, args["drug_id"])
    before = current_inventory_value(session, drug)
    if before is None:
        raise ValueError("inventory value is not numeric for this drug")
    qty = float(args["quantity"])
    mode = args["mode"]
    if mode == "decrease":
        after = before - qty
    elif mode == "increase":
        after = before + qty
    elif mode == "set":
        after = qty
    else:
        raise ValueError("unsupported mode")
    session["inventory_state"][drug["drug_id"]] = after
    payload = {
        "drug_id": drug["drug_id"],
        "mode": mode,
        "quantity": qty,
        "before": before,
        "after": after,
        "note": args.get("note", ""),
    }
    append_event(session, "update_inventory_balance", payload)
    _persist_session(session)
    inventory_update = {
        "drug_id": drug["drug_id"],
        "before": before,
        "after": after,
        "unit": drug.get("单位", ""),
        "mode": mode,
        "quantity": qty,
    }
    return result_with_session(session, {"inventory_update": inventory_update, **inventory_update})


def tool_update_record_status(args):
    session = get_session(args["session_id"])
    record_id = args["record_id"]
    record = session["records"].get(record_id)
    if not record:
        raise ValueError("unknown record_id")
    record["status"] = args["status"]
    if args.get("note"):
        record["status_note"] = args["note"]
    append_event(session, "update_record_status", {"record_id": record_id, "status": args["status"], "note": args.get("note", "")})
    _persist_session(session)
    return result_with_session(session, {"record": record, **record})


def tool_update_audit_decision(args):
    session = get_session(args["session_id"])
    record_id = args["record_id"]
    record = session["records"].get(record_id)
    if not record:
        raise ValueError("unknown record_id")
    if record["record_type"] not in {"audit_note", "insurance_intercept"}:
        raise ValueError("record_id is not an audit-like record")
    record["decision"] = args["decision"]
    if "reason" in args:
        record["reason"] = args.get("reason", "")
    if args.get("note"):
        record["decision_note"] = args["note"]
    append_event(session, "update_audit_decision", {
        "record_id": record_id,
        "decision": args["decision"],
        "reason": args.get("reason", ""),
        "note": args.get("note", ""),
    })
    _persist_session(session)
    return result_with_session(session, {"record": record, **record})


def tool_export_operation_log(args):
    session = get_session(args["session_id"])
    out = {
        "session_id": session["session_id"],
        "question_id": session["question_id"],
        "path": session["path"],
        "query_budget": session["query_budget"],
        "remaining_budget": session["remaining_budget"],
        "records": list(session["records"].values()),
        "inventory_state": session["inventory_state"],
        "events": session["events"],
        "log_path": str(session["log_path"]),
    }
    out_path = OUTPUT_DIR / f"{args['session_id']}.json"
    out_path.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    return result_with_session(session, {
        "export": {"export_path": str(out_path), **out},
        "export_path": str(out_path),
        **out,
    })


TOOLS = {
    "initialize_session": tool_create_session,
    "search_drug_catalog": tool_search_drug_catalog,
    "get_drug_base_info": tool_get_drug_base_info,
    "get_inventory_record": tool_get_inventory_record,
    "get_batch_record": tool_get_batch_record,
    "get_clinical_rule": tool_get_clinical_rule,
    "get_reimbursement_rule": tool_get_reimbursement_rule,
    "create_dispense_record": tool_create_dispense_record,
    "create_inventory_warning": tool_create_inventory_warning,
    "create_audit_note": tool_create_audit_note,
    "create_insurance_intercept": tool_create_insurance_intercept,
    "create_handoff_note": tool_create_handoff_note,
    "update_inventory_balance": tool_update_inventory_balance,
    "update_record_status": tool_update_record_status,
    "update_audit_decision": tool_update_audit_decision,
    "export_operation_log": tool_export_operation_log,
    "list_question_profiles": tool_list_question_profiles,
    # Aliases (legacy tool names used in some question profiles).
    "search_catalog": tool_search_drug_catalog,
    "get_inventory_position": tool_get_inventory_record,
    "get_batch_lots": tool_get_batch_record,
    "get_clinical_rules": tool_get_clinical_rule,
}


def ok(msg_id, result):
    return {"jsonrpc": "2.0", "id": msg_id, "result": result}


def err(msg_id, message):
    return {"jsonrpc": "2.0", "id": msg_id, "error": {"code": -32000, "message": message}}


def handle(req):
    method = req.get("method")
    msg_id = req.get("id")
    # Notifications have no "id" and require no response
    if msg_id is None and method and method.startswith("notifications/"):
        return None
    params = req.get("params") or {}
    if method == "initialize":
        return ok(msg_id, {
            "protocolVersion": "2024-11-05",
            "capabilities": {"tools": {}},
            "serverInfo": {"name": TOOL_SCHEMA["server_name"], "version": "0.3.0"},
        })
    if method == "session/create":
        try:
            result = initialize_question_session(params["question_id"], params.get("budget_override"))
        except Exception as e:
            return err(msg_id, str(e))
        return ok(msg_id, result)
    if method == "tools/list":
        tools = []
        for t in TOOL_SCHEMA["tools"]:
            entry = dict(t)
            if "input_schema" in entry and "inputSchema" not in entry:
                entry["inputSchema"] = entry.pop("input_schema")
            tools.append(entry)
        return ok(msg_id, {"tools": tools})
    if method == "tools/call":
        name = params["name"]
        args = params.get("arguments") or {}
        if name not in TOOLS:
            return err(msg_id, f"unknown tool: {name}")
        try:
            result = TOOLS[name](args)
        except Exception as e:
            return err(msg_id, str(e))
        return ok(msg_id, {"content": [{"type": "text", "text": json.dumps(result, ensure_ascii=False)}]})
    return err(msg_id, f"unsupported method: {method}")


def main():
    for raw in sys.stdin:
        raw = raw.strip()
        if not raw:
            continue
        try:
            req = json.loads(raw)
            resp = handle(req)
        except Exception as e:
            resp = err(None, f"invalid request: {e}")
        if resp is not None:
            sys.stdout.write(json.dumps(resp, ensure_ascii=False) + "\n")
            sys.stdout.flush()


if __name__ == "__main__":
    main()
